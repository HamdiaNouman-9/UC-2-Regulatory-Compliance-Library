"""
Text extraction for the Office-format attachments PDF-only Text_Extractor.py
never handled -- .docx (Word) and .xls/.xlsx (Excel). Regulator attachment
bundles (extra_meta.attachment_links) routinely mix these in with PDFs -- the
6687 draft-amendments bundle pairs a .pdf with a .docx comment form, and
several Ministry of Commerce/SDAIA rows carry .xlsx annexes -- and until now
dynamic_crawler.formfill.textinput.decide()'s fetch_file_text callback simply
returned None for anything that wasn't a PDF, silently dropping the attachment
rather than erroring.

No OCR here: both formats store text natively (unlike a scanned PDF), so this
is direct extraction, not the smart/scanned-detection path Text_Extractor.py
needs for PDFs.

    from processor.office_text_extractor import extract_office_text
    text = extract_office_text("/tmp/form.docx")   # by extension
"""

import logging
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


def extract_docx_text(path: str) -> str:
    """Paragraphs, in order, then every table's cells (regulator forms are
    often table-laid-out, e.g. a comment-submission form's field grid) --
    both are real content and dropping either loses information a PDF
    equivalent wouldn't."""
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_xlsx_text(path: str) -> str:
    """Every sheet, every non-empty row, cells joined with tabs -- close to
    what a person reading the spreadsheet would see, not a data model of it.
    data_only=True reads a formula's last-computed value rather than the
    formula text itself, since the value is what a compliance reader needs."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        parts = []
        for sheet in wb.worksheets:
            sheet_lines = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    sheet_lines.append("\t".join(cells))
            if sheet_lines:
                parts.append(f"=== Sheet: {sheet.title} ===")
                parts.extend(sheet_lines)
        return "\n".join(parts)
    finally:
        wb.close()


def extract_xls_text(path: str) -> str:
    """Legacy binary .xls -- a different file format from .xlsx entirely
    (OLE2, not a zip of XML), which openpyxl cannot open at all. xlrd 2.x
    dropped .xlsx support to focus solely on this, which is exactly the
    complementary role wanted here: openpyxl for .xlsx, xlrd for .xls."""
    import xlrd

    wb = xlrd.open_workbook(path)
    parts = []
    for sheet in wb.sheets():
        sheet_lines = []
        for row_idx in range(sheet.nrows):
            cells = [str(v).strip() for v in sheet.row_values(row_idx)
                    if str(v).strip()]
            if cells:
                sheet_lines.append("\t".join(cells))
        if sheet_lines:
            parts.append(f"=== Sheet: {sheet.name} ===")
            parts.extend(sheet_lines)
    return "\n".join(parts)


_EXTRACTORS = {
    ".docx": extract_docx_text,
    ".xlsx": extract_xlsx_text,
    ".xls": extract_xls_text,
}


def extract_office_text(path: str, suffix: Optional[str] = None) -> Optional[str]:
    """Dispatch by extension. Returns None (never raises) for an unsupported
    type or a file that fails to parse -- same contract as
    Text_Extractor.OCRProcessor's callers expect: a bad attachment must not
    take the rest of the document's real content down with it."""
    ext = (suffix or Path(path).suffix or "").lower()
    fn = _EXTRACTORS.get(ext)
    if fn is None:
        return None
    try:
        text = fn(path)
        return text.strip() if text else ""
    except Exception as e:
        logger.warning("office text extraction failed for %s (%s): %s", path, ext, e)
        return None


__all__ = ["extract_office_text", "extract_docx_text", "extract_xlsx_text", "extract_xls_text"]
