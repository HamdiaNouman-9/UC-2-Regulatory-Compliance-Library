import logging
import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Dict, Optional, Tuple

logger = logging.getLogger(__name__)

# Every extractor: (local file path) -> (text, metadata). metadata["method"] is
# what ends up in fs_file_text.method.


def extract_pdf(path: str) -> Tuple[str, Dict]:
    from processor.Text_Extractor import OCRProcessor
    text, meta = OCRProcessor.extract_text_from_pdf_smart(path)
    ocr = meta.get("ocr_pages") or 0
    kept = meta.get("good_pages") or 0
    meta["method"] = "native" if ocr == 0 else ("ocr" if ocr >= kept else "mixed")
    return text, meta


def extract_office(path: str) -> Tuple[str, Dict]:
    # .docx / .xlsx / .xls -- the extractor the pipeline already uses.
    from processor.office_text_extractor import extract_office_text
    return extract_office_text(path) or "", {"method": "office"}


def extract_pptx(path: str) -> Tuple[str, Dict]:
    from pptx import Presentation
    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if parts:
            slides.append(f"SLIDE {i}\n" + "\n".join(parts))
    return "\n\n".join(slides), {"method": "office", "total_pages": len(prs.slides)}


def _soffice() -> Optional[str]:
    default = r"C:\Program Files\LibreOffice\program\soffice.exe"
    return (os.getenv("SOFFICE_PATH") or shutil.which("soffice")
            or (default if os.path.exists(default) else None))


def extract_via_pdf(path: str) -> Tuple[str, Dict]:
    """Old binary .doc/.ppt: convert to PDF with LibreOffice, then read it like any
    PDF -- which also OCRs scanned images inside the file."""
    exe = _soffice()
    if not exe:
        raise RuntimeError("LibreOffice not found -- old .doc/.ppt need it. "
                           "Install it or set SOFFICE_PATH.")
    with tempfile.TemporaryDirectory() as out:
        subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", out, path],
                       check=True, timeout=180, capture_output=True)
        text, meta = extract_pdf(str(Path(out) / (Path(path).stem + ".pdf")))
    meta["method"] = "converted"
    return text, meta


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_office,
    ".xlsx": extract_office,
    ".xls": extract_office,
    ".pptx": extract_pptx,
    ".doc": extract_via_pdf,
    ".ppt": extract_via_pdf,
}


def extract(path: str, extension: str) -> Optional[Tuple[str, Dict]]:
    """None for a type we don't extract text from (e.g. .zip)."""
    fn = EXTRACTORS.get((extension or "").lower())
    return fn(path) if fn else None
